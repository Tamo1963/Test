"""Design report (.docx)."""
from __future__ import annotations

import datetime as dt
import io
import tempfile
import os

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor

from core.kin_curves import kinetics_table, monod_curves
from core.adm1 import IX


def _table(doc, rows, header=None, widths=None, style='Light Grid Accent 1'):
    t = doc.add_table(rows=0, cols=len(rows[0]) if rows else len(header))
    t.style = style
    if header:
        cells = t.add_row().cells
        for c, h in zip(cells, header):
            c.text = ''
            r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = ''
            c.paragraphs[0].add_run(str(v)).font.size = Pt(9)
    if widths:
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)
    return t


def _fig_to_doc(doc, fig, width_cm=15.5, caption=None):
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=170, bbox_inches='tight')
    plt.close(fig)
    buf.seek(0)
    doc.add_picture(buf, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)


def _f(x, n=2):
    return f'{x:,.{n}f}'


def build_report(design: dict, cad_png: str | None = None, scenarios=None, project='Anaerobic CSTR design') -> bytes:
    di, fp, geo = design['inputs'], design['feed'], design['geometry']
    perf, mix, heat, coil, en = design['performance'], design['mixing'], design['heat'], design['coil'], design['energy']
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)

    t = doc.add_heading(project, 0)
    doc.add_paragraph(f'Generated {dt.date.today().isoformat()} by adm_cstr_designer (ADM1-based preliminary design)')
    p = doc.add_paragraph()
    r = p.add_run('Preliminary design. Feedstock properties and several kinetic constants are indicative defaults; '
                  'verify with laboratory analysis and BMP tests before detailed design.')
    r.bold = True; r.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)

    # 1. summary ---------------------------------------------------------
    doc.add_heading('1. Design summary', 1)
    _table(doc, [
        ['Working volume', f"{_f(design['spec'].V_liq, 0)} m3 ({geo['n_tanks']} x {_f(geo['V_liq_each'], 0)} m3)"],
        ['Tank', f"D = {_f(geo['D'])} m, liquid depth {_f(geo['H_liq'])} m, total height {_f(geo['H_total'])} m"],
        ['HRT / OLR', f"{_f(di['HRT_d'], 1)} d / {_f(perf['OLR_kgVS_m3_d'])} kg VS/m3/d"],
        ['Operating temperature', f"{_f(di['T_op_C'], 1)} C"],
        ['Methane / biogas', f"{_f(perf['CH4_Nm3_d'], 0)} / {_f(perf['biogas_Nm3_d'], 0)} Nm3/d ({_f(perf['CH4_pct'], 1)} % CH4)"],
        ['Specific methane yield', f"{_f(perf['SMY_Nm3_kgVS'], 3)} Nm3/kg VS fed (blend BMP {_f(perf['BMP_blend'], 3)})"],
        ['Effluent pH / VFA', f"{_f(perf['pH'])} / {_f(perf['VFA_kgCOD_m3'])} kg COD/m3"],
        ['Agitators', f"{mix['n_agitators']} x {_f(mix['motor_kW'], 1)} kW motor, D = {_f(mix['D_impeller_m'])} m, {_f(mix['N_rpm'], 1)} rpm"],
        ['Heating coil', f"{_f(coil['area_m2'], 1)} m2, {_f(coil['length_m'], 0)} m of DN{50} pipe, duty {_f(heat['Q_design_kW'], 0)} kW per tank"],
        ['Net energy (biogas - heat - mixing)', f"{_f(en['net_kW'], 0)} kW thermal-equivalent ({_f(en['parasitic_pct'], 1)} % parasitic)"],
    ], header=['Item', 'Result'], widths=[5, 11])
    if design['flags']:
        doc.add_paragraph()
        doc.add_paragraph('Process flags:', style='Intense Quote')
        for fl in design['flags']:
            doc.add_paragraph(fl, style='List Bullet')

    # 2. inputs ---------------------------------------------------------------------
    doc.add_heading('2. Inputs and design basis', 1)
    _table(doc, [
        ['Substrate flow (wet)', f"{_f(di['Q_tpd'], 1)} t/d  = {_f(fp['Q_m3d'], 1)} m3/d (density {_f(fp['rho'], 0)} kg/m3)"],
        ['Feed temperature', f"{_f(di['T_feed_C'], 1)} C"],
        ['TS (wet basis) / VS (of TS)', f"{_f(di['TS_pct'], 1)} % / {_f(di['VS_pct_TS'], 1)} %"],
        ['VS load', f"{_f(fp['VS_load_kg_d'], 0)} kg VS/d  ({_f(fp['VS_kg_m3'], 1)} kg VS/m3)"],
        ['Feed COD', f"{_f(fp['COD_kg_m3'], 1)} kg COD/m3 (blend {_f(fp['COD_per_VS'])} kg COD/kg VS)"],
        ['HRT / T operating / T ambient', f"{di['HRT_d']} d / {di['T_op_C']} C / {di['T_amb_C']} C"],
    ], header=['Input', 'Value'], widths=[5, 11])
    doc.add_paragraph()
    doc.add_paragraph('Feedstock blend (share of VS) and library properties')
    rows = []
    for k, w in fp['weights'].items():
        rows.append([k, f'{100 * w:.1f} %'])
    _table(doc, rows, header=['Feedstock', 'Share of VS'], widths=[8, 4])
    fr = fp['fractions']
    doc.add_paragraph(
        f"Biodegradable COD split in the ADM1 feed: carbohydrate {100 * fr['ch']:.1f} %, protein {100 * fr['pr']:.1f} %, "
        f"lipid {100 * fr['li']:.1f} %, acetate {100 * fr['vfa']:.1f} %; inert/slowly-degradable {100 * fr['inert']:.1f} %. "
        f"The inert fraction is set so the blend's maximum methane equals its library BMP. Blend hydrolysis constants: "
        f"k_ch = {fp['hyd']['k_hyd_ch']:.2f}, k_pr = {fp['hyd']['k_hyd_pr']:.2f}, k_li = {fp['hyd']['k_hyd_li']:.2f} 1/d.")

    # 3. kinetics -----------------------------------------------------------------------
    doc.add_heading('3. Kinetics (Monod / ADM1)', 1)
    ks = design['kinetics']
    doc.add_paragraph(f"Constant set: {design['inputs']['kinetics_set']}. Source: {ks['reference']}")
    kt = kinetics_table(ks, di['T_op_C'], di['theta'])
    cols = ['Group', 'k_m @ref (1/d)', f"k_m @ {di['T_op_C']:g} C (1/d)", 'K_S (kg COD/m3)', 'Y (kg/kg)', 'b (1/d)']
    _table(doc, [[r['Group']] + [f'{r[c]:.4g}' for c in cols[1:]] for _, r in kt.iterrows()], header=cols)
    doc.add_paragraph(
        'Rate form: r = k_m * S/(K_S + S) * X * I_pH * I_IN * I_H2 * I_NH3 (inhibition factors apply to the relevant groups). '
        'Temperature correction on biological rates is theta^(T - Tref) with theta from the set; acid-base and gas-solubility '
        'constants use van t Hoff corrections.')
    fig, ax = plt.subplots(figsize=(6.5, 3.4))
    for lab, (S, r) in monod_curves(ks, di['T_op_C'], di['theta']).items():
        ax.plot(S / S.max(), r, label=lab)
    ax.set_xlabel('S / (5 K_S)'); ax.set_ylabel('specific uptake rate (kg COD/kg COD/d)')
    ax.legend(fontsize=6, ncol=2); ax.grid(alpha=0.3)
    _fig_to_doc(doc, fig, caption='Figure 1. Monod uptake curves of the seven ADM1 microbial groups at the operating temperature.')

    # 4. process results ----------------------------------------------------------------
    doc.add_heading('4. Process performance (steady state)', 1)
    s = design['state']
    _table(doc, [
        ['pH', _f(perf['pH'])],
        ['Total VFA (as COD)', f"{_f(perf['VFA_kgCOD_m3'], 3)} kg/m3"],
        ['Inorganic N / free NH3-N', f"{_f(perf['S_IN_mgN_L'], 0)} / {_f(perf['NH3_mgN_L'], 0)} mg N/L"],
        ['Biogas / CH4 / CO2', f"{_f(perf['biogas_Nm3_d'], 0)} / {_f(perf['CH4_Nm3_d'], 0)} / {_f(perf['CO2_Nm3_d'], 0)} Nm3/d"],
        ['CH4 content', f"{_f(perf['CH4_pct'], 1)} %"],
        ['Specific methane yield', f"{_f(perf['SMY_Nm3_kgVS'], 3)} Nm3/kg VS (conversion {_f(perf['conversion_pct'], 0)} % of BMP)"],
        ['COD removal', f"{_f(perf['COD_removal_pct'], 1)} % (COD balance closure {_f(100 * perf['cod_closure'], 1)} %)"],
        ['Methane energy', f"{_f(perf['CH4_energy_kW'], 0)} kW (LHV 35.8 MJ/Nm3)"],
        ['Acetoclastic / hydrogenotrophic biomass', f"{_f(perf['X_ac'], 2)} / {_f(perf['X_h2'], 2)} kg COD/m3"],
    ], header=['Quantity', 'Value'], widths=[6, 10])
    h = design.get('history')
    if h is not None:
        fig, ax = plt.subplots(1, 3, figsize=(10, 3))
        ax[0].plot(h['t'], h['ch4']); ax[0].set_title('CH4 (Nm3/d)')
        ax[1].plot(h['t'], h['pH']); ax[1].set_title('pH')
        ax[2].plot(h['t'], h['vfa']); ax[2].set_title('VFA (kg COD/m3)')
        for a in ax:
            a.set_xlabel('time (d)'); a.grid(alpha=0.3)
        fig.tight_layout()
        _fig_to_doc(doc, fig, 16, 'Figure 2. Start-up trajectory from a seeded digester to steady state.')

    # 5. agitator --------------------------------------------------------------------------
    doc.add_heading('5. Agitator sizing', 1)
    doc.add_paragraph(
        f"Method: top-entry axial-flow hydrofoil; specific power input {mix['eps_W_m3']} W/m3; apparent viscosity "
        f"{mix['mu_Pa_s']:.3f} Pa.s (indicative function of TS); power number {mix['Np']}, flow number {mix['Nq']}, "
        f"D/T = {mix['D_over_T']:.2f}. N = (P/(Np rho D^5))^(1/3).")
    _table(doc, [
        ['Number of agitators', mix['n_agitators']], ['Impeller diameter', f"{_f(mix['D_impeller_m'])} m"],
        ['Speed', f"{_f(mix['N_rpm'], 1)} rpm"], ['Reynolds number', f"{mix['Re']:,.0f}"],
        ['Tip speed', f"{_f(mix['tip_speed_m_s'])} m/s"], ['Pumping capacity (each)', f"{mix['pumping_m3_h']:,.0f} m3/h"],
        ['Turnover time', f"{_f(mix['turnover_min'], 1)} min"], ['Mean velocity gradient G', f"{_f(mix['G_1_s'], 1)} 1/s"],
        ['Hydraulic power (each)', f"{_f(mix['P_hydraulic_kW'])} kW"], ['Shaft power (each)', f"{_f(mix['P_shaft_kW'])} kW"],
        ['Motor (IEC, with 15 % service factor)', f"{mix['motor_kW']} kW"],
        ['Torque / shaft diameter (peak 2.5x, tau 40 MPa)', f"{mix['torque_Nm']:,.0f} N.m / {_f(mix['shaft_diameter_mm'], 0)} mm"],
    ], header=['Parameter', 'Value'], widths=[8, 8])
    for fl in mix['flags']:
        doc.add_paragraph(fl, style='List Bullet')

    # 6. heating ---------------------------------------------------------------------------
    doc.add_heading('6. Heat load and heating coil', 1)
    _table(doc, [
        ['Feed heating', f"{_f(heat['Q_feed_kW'], 1)} kW"], ['Wall loss', f"{_f(heat['Q_wall_kW'], 1)} kW"],
        ['Roof loss', f"{_f(heat['Q_roof_kW'], 1)} kW"], ['Floor loss', f"{_f(heat['Q_floor_kW'], 1)} kW"],
        ['Design duty (+15 % margin)', f"{_f(heat['Q_design_kW'], 1)} kW per tank"],
    ], header=['Component', 'Value'], widths=[8, 8])
    fig, ax = plt.subplots(figsize=(5, 2.4))
    ax.barh(['Feed heating', 'Wall', 'Roof', 'Floor'],
            [heat['Q_feed_kW'], heat['Q_wall_kW'], heat['Q_roof_kW'], heat['Q_floor_kW']], color='#d9534f')
    ax.set_xlabel('kW'); ax.grid(alpha=0.3, axis='x')
    _fig_to_doc(doc, fig, 11, 'Figure 3. Heat-load breakdown per tank.')
    doc.add_paragraph(
        f"Coil: internal wall-mounted helix, hot water {coil['T_w_in']:.0f}/{coil['T_w_out']:.0f} C inside, digestate at "
        f"{di['T_op_C']:.0f} C outside. U = {coil['U_W_m2K']:.0f} W/m2K (indicative; decreases with TS), LMTD = {coil['LMTD_K']:.1f} K.")
    _table(doc, [
        ['Required area', f"{_f(coil['area_m2'], 1)} m2"], ['Pipe', f"OD {coil['OD_mm']:.1f} / ID {coil['ID_mm']:.1f} mm (DN50)"],
        ['Total length', f"{_f(coil['length_m'], 0)} m in {coil['n_circuits']} circuit(s) of {_f(coil['length_per_circuit_m'], 0)} m"],
        ['Water flow', f"{_f(coil['water_flow_m3_h'], 2)} m3/h, velocity {_f(coil['water_velocity_m_s'], 2)} m/s"],
        ['Coil diameter / turns / pitch', f"{_f(coil['coil_diameter_m'], 2)} m / {_f(coil['turns_total'], 1)} / {_f(coil['pitch_m'], 2)} m"],
    ], header=['Parameter', 'Value'], widths=[6, 10])
    for fl in coil['flags']:
        doc.add_paragraph(fl, style='List Bullet')

    # 7. energy ---------------------------------------------------------------------------
    doc.add_heading('7. Energy balance', 1)
    _table(doc, [['Biogas (methane LHV)', f"{_f(en['biogas_kW'], 0)} kW"], ['Digester heating', f"{_f(en['heat_kW'], 0)} kW"],
                 ['Mixing (electrical, input)', f"{_f(en['mixing_kW'], 1)} kW"], ['Net', f"{_f(en['net_kW'], 0)} kW"],
                 ['Parasitic share', f"{_f(en['parasitic_pct'], 1)} %"]], header=['Item', 'Value'], widths=[8, 8])
    doc.add_paragraph('Heating is a thermal duty and mixing is electrical; the net figure is a simple energy-equivalent comparison, not a CHP model.')

    # 8. CAD ---------------------------------------------------------------------------------
    if cad_png and os.path.exists(cad_png):
        doc.add_heading('8. 3D model', 1)
        doc.add_picture(cad_png, width=Cm(11))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph('Figure 4. Tank, heating coil and agitators (preview render; the FreeCAD model uses the same dimensions).')

    # 9. scenarios --------------------------------------------------------------------------
    if scenarios is not None and len(scenarios):
        doc.add_heading('9. Scenario comparison', 1)
        cols = [c for c in scenarios.columns]
        _table(doc, [[str(r[c]) for c in cols] for _, r in scenarios.iterrows()], header=cols)

    # 10. validation + limitations ----------------------------------------------------
    doc.add_heading('10. Model validation and limitations', 1)
    for line in [
        'Validation: the ADM1 implementation reproduces the BSM2 benchmark steady state (Q = 178.5 m3/d, V = 3,400 m3, 35 C) '
        'within 1 % on all states (published pH 7.26 reproduced; COD balance closes within 1 %). See tests/test_bsm2_validation.py.',
        'Feedstock properties (TS, VS, COD/VS, BMP, composition, hydrolysis constants, ammonia release) are indicative '
        'literature-range defaults, not measured values. Measure your own material.',
        'ADM1 is a homogeneous model: particle-size effects, floating layers, stratification, and trace-element or sulphide '
        'effects are not represented. ADM1 does not wash out methanogens until HRT is roughly 3 d, so short-HRT results are optimistic.',
        'Temperature correction of biological rates is an extrapolation outside about 30-42 C. Thermophilic design needs a calibrated set.',
        'Viscosity, wall U-values and coil U are indicative correlations; confirm with rheometry and vendor data.',
        'Agitator sizing uses specific power and impeller flow/power numbers; confirm with the agitator vendor (CFD) for final selection.',
    ]:
        doc.add_paragraph(line, style='List Bullet')

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
